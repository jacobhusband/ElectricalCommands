import os
import sys
import tempfile
import threading
import types
import unittest
from pathlib import Path
from unittest.mock import patch


def _ensure_google_genai_stub():
    try:
        from google import genai as _genai  # noqa: F401
        from google.genai import types as _types  # noqa: F401
        return
    except Exception:
        google_module = sys.modules.get("google")
        if google_module is None:
            google_module = types.ModuleType("google")
            google_module.__path__ = []
            sys.modules["google"] = google_module

        genai_module = types.ModuleType("google.genai")
        genai_types_module = types.ModuleType("google.genai.types")
        genai_module.types = genai_types_module
        google_module.genai = genai_module

        sys.modules["google.genai"] = genai_module
        sys.modules["google.genai.types"] = genai_types_module


def _ensure_webview_stub():
    try:
        import webview  # noqa: F401
        return
    except Exception:
        webview_module = types.ModuleType("webview")
        webview_module.windows = []
        webview_module.FOLDER_DIALOG = "folder"
        webview_module.OPEN_DIALOG = "open"
        webview_module.FileDialog = types.SimpleNamespace(SAVE="save")
        webview_module.create_window = lambda *args, **kwargs: None
        webview_module.start = lambda *args, **kwargs: None
        sys.modules["webview"] = webview_module


def _ensure_dotenv_stub():
    try:
        from dotenv import load_dotenv as _load_dotenv  # noqa: F401
        return
    except Exception:
        dotenv_module = types.ModuleType("dotenv")
        dotenv_module.load_dotenv = lambda *args, **kwargs: False
        sys.modules["dotenv"] = dotenv_module


def _ensure_requests_stub():
    try:
        import requests  # noqa: F401
        return
    except Exception:
        requests_module = types.ModuleType("requests")
        requests_module.get = lambda *args, **kwargs: None
        sys.modules["requests"] = requests_module


def _ensure_pydantic_stub():
    try:
        from pydantic import BaseModel as _BaseModel, Field as _Field  # noqa: F401
        return
    except Exception:
        pydantic_module = types.ModuleType("pydantic")

        class BaseModel:
            pass

        def Field(*args, **kwargs):
            if args:
                return args[0]
            return kwargs.get("default")

        pydantic_module.BaseModel = BaseModel
        pydantic_module.Field = Field
        sys.modules["pydantic"] = pydantic_module


def _ensure_pil_stub():
    try:
        from PIL import Image as _image  # noqa: F401
        return
    except Exception:
        pil_module = types.ModuleType("PIL")
        pil_image_module = types.ModuleType("PIL.Image")
        pil_module.Image = pil_image_module
        sys.modules["PIL"] = pil_module
        sys.modules["PIL.Image"] = pil_image_module


def _ensure_openpyxl_stub():
    try:
        import openpyxl  # noqa: F401
        from openpyxl.worksheet.copier import WorksheetCopy as _WorksheetCopy  # noqa: F401
        return
    except Exception:
        openpyxl_module = types.ModuleType("openpyxl")
        worksheet_module = types.ModuleType("openpyxl.worksheet")
        copier_module = types.ModuleType("openpyxl.worksheet.copier")

        class WorksheetCopy:
            pass

        copier_module.WorksheetCopy = WorksheetCopy
        worksheet_module.copier = copier_module
        openpyxl_module.worksheet = worksheet_module

        sys.modules["openpyxl"] = openpyxl_module
        sys.modules["openpyxl.worksheet"] = worksheet_module
        sys.modules["openpyxl.worksheet.copier"] = copier_module


_ensure_google_genai_stub()
_ensure_webview_stub()
_ensure_dotenv_stub()
_ensure_requests_stub()
_ensure_pydantic_stub()
_ensure_pil_stub()
_ensure_openpyxl_stub()

from main import Api
import main


class PanelScheduleAiBackendTests(unittest.TestCase):
    def setUp(self):
        self.api = Api.__new__(Api)

    def _create_panel_schedule_workbook(self, path):
        if not hasattr(main.openpyxl, "Workbook"):
            self.skipTest("openpyxl Workbook is not available")
        workbook = main.openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "TEMPLATE"
        worksheet["N3"] = "EXISTING"
        workbook.save(path)
        workbook.close()

    def _panel_data(self, circuits):
        return types.SimpleNamespace(
            panel_name="MDP",
            voltage="120/208V",
            bus_rating="100A",
            aic_rating="",
            phase="3",
            wire="4",
            mounting="SURFACE",
            enclosure="NEMA 1",
            circuits=circuits,
        )

    def test_process_panel_schedule_payload_uses_plural_path_lists(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            output_path = Path(temp_dir) / "panel_schedule.xlsx"
            output_path.write_text("placeholder", encoding="utf-8")
            captured = {}
            fake_panel_data = types.SimpleNamespace(panel_name="MDP")

            def fake_analyze(panel_name, breaker_paths, directory_paths, input_mode="field_photos"):
                captured["panel_name"] = panel_name
                captured["breaker_paths"] = list(breaker_paths)
                captured["directory_paths"] = list(directory_paths)
                captured["input_mode"] = input_mode
                return fake_panel_data

            with patch.object(
                self.api,
                "_analyze_panel_schedule_images",
                side_effect=fake_analyze,
            ), patch.object(
                self.api,
                "_update_panel_schedule_workbook",
                return_value="MDP",
            ) as update_mock:
                result = self.api._process_panel_schedule_payload(
                    {
                        "outputMode": "existing",
                        "outputPath": str(output_path),
                        "panels": [
                            {
                                "panelId": "panel_1",
                                "panelName": "MDP",
                                "breakerPaths": [
                                    "breaker-top.jpg",
                                    "",
                                    "breaker-bottom.jpg",
                                ],
                                "directoryPaths": [
                                    "directory-1-42.jpg",
                                    "directory-43-84.jpg",
                                ],
                            }
                        ],
                    }
                )

            self.assertEqual("success", result["status"])
            self.assertEqual("MDP", captured["panel_name"])
            self.assertEqual(
                ["breaker-top.jpg", "breaker-bottom.jpg"],
                captured["breaker_paths"],
            )
            self.assertEqual(
                ["directory-1-42.jpg", "directory-43-84.jpg"],
                captured["directory_paths"],
            )
            self.assertEqual("field_photos", captured["input_mode"])
            update_mock.assert_called_once_with(
                fake_panel_data,
                os.path.normpath(str(output_path)),
                use_extracted_kva=False,
            )

    def test_process_panel_schedule_payload_requires_at_least_one_breaker_and_directory_photo(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            output_path = Path(temp_dir) / "panel_schedule.xlsx"
            output_path.write_text("placeholder", encoding="utf-8")

            result = self.api._process_panel_schedule_payload(
                {
                    "outputMode": "existing",
                    "outputPath": str(output_path),
                    "panelName": "MDP",
                    "breakerPaths": [],
                    "directoryPaths": [],
                }
            )

        self.assertEqual("error", result["status"])
        self.assertEqual(0, result["successCount"])
        self.assertEqual(1, result["failureCount"])
        self.assertIn(
            "At least one breaker photo and at least one directory photo are required.",
            result["message"],
        )

    def test_process_panel_schedule_payload_existing_directory_requires_directory_photo(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            output_path = Path(temp_dir) / "panel_schedule.xlsx"
            output_path.write_text("placeholder", encoding="utf-8")

            result = self.api._process_panel_schedule_payload(
                {
                    "outputMode": "existing",
                    "outputPath": str(output_path),
                    "panelName": "MDP",
                    "inputMode": "existing_directory",
                    "breakerPaths": [],
                    "directoryPaths": [],
                }
            )

        self.assertEqual("error", result["status"])
        self.assertEqual(0, result["successCount"])
        self.assertEqual(1, result["failureCount"])
        self.assertIn(
            "At least one directory photo is required for existing directory mode.",
            result["message"],
        )

    def test_process_panel_schedule_payload_existing_directory_success_with_only_directory(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            output_path = Path(temp_dir) / "panel_schedule.xlsx"
            output_path.write_text("placeholder", encoding="utf-8")
            fake_panel_data = types.SimpleNamespace(panel_name="MDP")

            with patch.object(
                self.api,
                "_analyze_panel_schedule_images",
                return_value=fake_panel_data,
            ) as analyze_mock, patch.object(
                self.api,
                "_update_panel_schedule_workbook",
                return_value="MDP",
            ) as update_mock:
                result = self.api._process_panel_schedule_payload(
                    {
                        "outputMode": "existing",
                        "outputPath": str(output_path),
                        "panelName": "MDP",
                        "inputMode": "existing_directory",
                        "breakerPaths": [],
                        "directoryPaths": ["directory.jpg"],
                    }
                )

        self.assertEqual("success", result["status"])
        analyze_mock.assert_called_once_with(
            "MDP", [], ["directory.jpg"], input_mode="existing_directory"
        )
        update_mock.assert_called_once_with(
            fake_panel_data,
            os.path.normpath(str(output_path)),
            use_extracted_kva=True,
        )

    def test_build_panel_schedule_prompt_mentions_split_breakers_and_directories(self):
        prompt = self.api._build_panel_schedule_prompt("MDP", 3, 2)

        self.assertIn("upper half, middle, and bottom half images", prompt)
        self.assertIn("circuits 1-42 on one image and 43-84 on another", prompt)
        self.assertIn("Do not treat split photos as separate panels.", prompt)
        self.assertIn(
            "Use visible breaker positions and circuit numbering to merge split sections",
            prompt,
        )
        self.assertIn('JSON field "aic_rating"', prompt)
        self.assertIn("Do not confuse AIC rating with Bus Rating", prompt)

    def test_build_existing_directory_prompt_requires_visible_kva_extraction(self):
        prompt = self.api._build_existing_directory_prompt("MDP", 2)

        self.assertIn('JSON field "kva"', prompt)
        self.assertIn('JSON field "phase_kva"', prompt)
        self.assertIn("visible kVA decimal", prompt)
        self.assertIn("Do not add, total, or sum phase loads together", prompt)
        self.assertIn("Do not repeat a 2-pole or 3-pole total", prompt)
        self.assertIn("Do not estimate kVA", prompt)
        self.assertIn("do not infer it from breaker amperage", prompt)
        self.assertIn('JSON field "aic_rating"', prompt)
        self.assertIn("Do not confuse AIC rating with Bus Rating", prompt)

    def test_workbook_writes_extracted_aic_rating_to_header(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([])
            panel_data.aic_rating = "22,000 AIC"

            sheet_name = main.cb_update_excel_workbook(
                panel_data,
                str(workbook_path),
                use_extracted_kva=True,
            )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual("22,000 AIC", worksheet["N3"].value)
            finally:
                workbook.close()

    def test_workbook_preserves_template_aic_rating_when_extraction_blank(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([])
            panel_data.aic_rating = ""

            sheet_name = main.cb_update_excel_workbook(
                panel_data,
                str(workbook_path),
                use_extracted_kva=True,
            )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual("EXISTING", worksheet["N3"].value)
            finally:
                workbook.close()

    def test_existing_directory_workbook_uses_phase_kva_without_estimator(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([
                types.SimpleNamespace(
                    circuit_number=1,
                    description="LIGHTING",
                    breaker_amps="20",
                    poles=2,
                    load_type="C",
                    kva="2.30 kVA",
                    phase_kva=["1.20 kVA", "1.10 kVA"],
                )
            ])

            with patch(
                "main.cb_calculate_estimated_load",
                side_effect=AssertionError("estimator should not run"),
            ):
                sheet_name = main.cb_update_excel_workbook(
                    panel_data,
                    str(workbook_path),
                    use_extracted_kva=True,
                )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual(1.2, worksheet["I8"].value)
                self.assertEqual(1.1, worksheet["I9"].value)
            finally:
                workbook.close()

    def test_existing_directory_workbook_writes_three_pole_phase_kva(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([
                types.SimpleNamespace(
                    circuit_number=1,
                    description="RTU",
                    breaker_amps="60",
                    poles=3,
                    load_type="M",
                    kva="3.30 kVA",
                    phase_kva=["0.90", "1.10", "1.30"],
                )
            ])

            with patch(
                "main.cb_calculate_estimated_load",
                side_effect=AssertionError("estimator should not run"),
            ):
                sheet_name = main.cb_update_excel_workbook(
                    panel_data,
                    str(workbook_path),
                    use_extracted_kva=True,
                )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual(0.9, worksheet["I8"].value)
                self.assertEqual(1.1, worksheet["I9"].value)
                self.assertEqual(1.3, worksheet["I10"].value)
            finally:
                workbook.close()

    def test_existing_directory_workbook_legacy_kva_fallback_only_top_row(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([
                types.SimpleNamespace(
                    circuit_number=1,
                    description="PUMP",
                    breaker_amps="30",
                    poles=2,
                    load_type="M",
                    kva="1.50 kVA",
                )
            ])

            with patch(
                "main.cb_calculate_estimated_load",
                side_effect=AssertionError("estimator should not run"),
            ):
                sheet_name = main.cb_update_excel_workbook(
                    panel_data,
                    str(workbook_path),
                    use_extracted_kva=True,
                )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual(1.5, worksheet["I8"].value)
                self.assertIn(worksheet["I9"].value, ("", None))
            finally:
                workbook.close()

    def test_existing_directory_workbook_leaves_missing_extracted_kva_blank(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([
                types.SimpleNamespace(
                    circuit_number=2,
                    description="RECEPTACLES",
                    breaker_amps="20",
                    poles=1,
                    load_type="G",
                    kva="",
                )
            ])

            with patch(
                "main.cb_calculate_estimated_load",
                side_effect=AssertionError("estimator should not run"),
            ):
                sheet_name = main.cb_update_excel_workbook(
                    panel_data,
                    str(workbook_path),
                    use_extracted_kva=True,
                )

            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertIn(worksheet["K8"].value, ("", None))
            finally:
                workbook.close()

    def test_field_photo_workbook_still_uses_estimator(self):
        with tempfile.TemporaryDirectory(prefix="acies-panel-schedule-") as temp_dir:
            workbook_path = Path(temp_dir) / "panel_schedule.xlsx"
            self._create_panel_schedule_workbook(workbook_path)
            panel_data = self._panel_data([
                types.SimpleNamespace(
                    circuit_number=1,
                    description="LIGHTING",
                    breaker_amps="20",
                    poles=1,
                    load_type="C",
                    kva="9.99",
                )
            ])

            with patch("main.cb_calculate_estimated_load", return_value=0.42) as estimator:
                sheet_name = main.cb_update_excel_workbook(
                    panel_data,
                    str(workbook_path),
                    use_extracted_kva=False,
                )

            estimator.assert_called_once_with("20", "LIGHTING", "C")
            workbook = main.openpyxl.load_workbook(workbook_path)
            try:
                worksheet = workbook[sheet_name]
                self.assertEqual(0.42, worksheet["I8"].value)
            finally:
                workbook.close()

    def test_register_heif_support_calls_register_opener_when_available(self):
        calls = []
        fake_module = types.SimpleNamespace(
            register_heif_opener=lambda: calls.append("registered")
        )

        with patch.object(main, "HEIF_SUPPORT_ENABLED", False), patch.object(
            main,
            "HEIF_SUPPORT_ERROR",
            "",
        ), patch("main.importlib.import_module", return_value=fake_module) as import_mock:
            result = main._register_heif_support()
            self.assertTrue(main.HEIF_SUPPORT_ENABLED)
            self.assertEqual("", main.HEIF_SUPPORT_ERROR)

        self.assertTrue(result)
        self.assertEqual(["registered"], calls)
        import_mock.assert_called_once_with("pillow_heif")

    def test_register_heif_support_degrades_cleanly_when_import_fails(self):
        with patch.object(main, "HEIF_SUPPORT_ENABLED", True), patch.object(
            main,
            "HEIF_SUPPORT_ERROR",
            "",
        ), patch(
            "main.importlib.import_module",
            side_effect=ModuleNotFoundError("No module named 'pillow_heif'"),
        ):
            result = main._register_heif_support()
            self.assertFalse(main.HEIF_SUPPORT_ENABLED)
            self.assertIn("pillow_heif", main.HEIF_SUPPORT_ERROR)

        self.assertFalse(result)

    def test_open_local_pil_image_maps_unreadable_heic_to_clear_value_error(self):
        with patch.object(main, "HEIF_SUPPORT_ENABLED", True), patch(
            "main.PILImage.open",
            side_effect=main.UnidentifiedImageError("cannot identify image file"),
        ):
            with self.assertRaises(ValueError) as ctx:
                main._open_local_pil_image("panel_photo.HEIC")

        self.assertIn("HEIC/HEIF", str(ctx.exception))
        self.assertNotIn("cannot identify image file", str(ctx.exception))

    def test_analyze_panel_schedule_images_uses_shared_heic_image_opener(self):
        opened_paths = []
        fake_images = []
        generate_content_calls = []

        class FakeImage:
            def __init__(self, path):
                self.path = path
                self.closed = False

            def close(self):
                self.closed = True

            def thumbnail(self, size, resampling=None):
                self.thumbnail_size = size

        def fake_open(path):
            opened_paths.append(path)
            image = FakeImage(path)
            fake_images.append(image)
            return image

        class FakeModels:
            def generate_content(self, **kwargs):
                generate_content_calls.append(kwargs)
                return types.SimpleNamespace(parsed=types.SimpleNamespace(panel_name=""), text="")

        fake_client = types.SimpleNamespace(models=FakeModels())

        with patch.object(
            self.api,
            "_resolve_panel_schedule_api_key",
            return_value="test-key",
        ), patch.object(
            self.api,
            "_ensure_aiohttp",
            return_value=None,
        ), patch(
            "main.genai.Client",
            return_value=fake_client,
            create=True,
        ), patch(
            "main._open_local_pil_image",
            side_effect=fake_open,
        ), patch(
            "main.cb_enforce_rate_limit",
            return_value=None,
        ), patch(
            "main.types.GenerateContentConfig",
            side_effect=lambda **kwargs: kwargs,
            create=True,
        ), patch(
            "main.os.path.exists",
            return_value=True,
        ):
            result = self.api._analyze_panel_schedule_images(
                "MDP",
                ["breaker_top.HEIC"],
                ["directory_bottom.HEIC"],
            )

        self.assertEqual("MDP", result.panel_name)
        self.assertEqual(
            ["breaker_top.HEIC", "directory_bottom.HEIC"],
            opened_paths,
        )
        self.assertEqual("gemini-3.6-flash", generate_content_calls[0]["model"])
        self.assertTrue(all(image.closed for image in fake_images))

    def test_run_panel_schedule_background_returns_job_id_and_seeds_running_status(self):
        started = {}

        class FakeThread:
            def __init__(self, target=None, args=(), daemon=None):
                started["target"] = target
                started["args"] = args
                started["daemon"] = daemon

            def start(self):
                started["called"] = True

        payload = {
            "outputMode": "new",
            "outputPath": r"C:\temp\panel_schedule.xlsx",
            "activityId": "activity_panel_schedule_1",
            "panels": [
                {"panelId": "panel_1", "panelName": "MDP"},
                {"panelId": "panel_2", "panelName": "L1"},
            ],
        }

        with patch("main.threading.Thread", FakeThread):
            result = self.api.run_panel_schedule_background(payload)

        self.assertEqual("started", result["status"])
        self.assertTrue(result["jobId"])
        self.assertEqual(2, result["panelCount"])
        self.assertEqual("activity_panel_schedule_1", result["activityId"])
        self.assertTrue(started.get("called"))
        self.assertEqual(self.api._panel_schedule_worker, started["target"])
        self.assertEqual((result["jobId"], payload), started["args"])
        self.assertTrue(started["daemon"])

        status = self.api.get_panel_schedule_background_status(result["jobId"])
        self.assertEqual("running", status["status"])
        self.assertEqual(result["jobId"], status["jobId"])
        self.assertEqual(2, status["panelCount"])
        self.assertEqual(0, status["completedCount"])
        self.assertEqual("activity_panel_schedule_1", status["activityId"])
        self.assertEqual(os.path.normpath(payload["outputPath"]), status["outputPath"])
        self.assertEqual("", status["finishedAt"])
        self.assertTrue(status["startedAt"])

    def test_panel_schedule_worker_stores_terminal_status_when_js_notification_fails(self):
        payload = {
            "outputMode": "existing",
            "outputPath": r"C:\temp\panel_schedule.xlsx",
        }
        job_id = "job_panel_success"
        self.api._store_panel_schedule_job_record(
            self.api._build_panel_schedule_job_record(job_id, payload, 1)
        )
        self.api._panel_schedule_running = True

        def raise_evaluate_js(_script):
            raise RuntimeError("js bridge offline")

        fake_window = types.SimpleNamespace(evaluate_js=raise_evaluate_js)

        with patch.object(
            self.api,
            "_process_panel_schedule_payload",
            return_value={
                "status": "success",
                "message": "Added 1 panel sheet to schedule.",
                "outputPath": os.path.normpath(payload["outputPath"]),
                "outputFolder": r"C:\temp",
                "successCount": 1,
                "failureCount": 0,
                "results": [
                    {
                        "panelId": "panel_1",
                        "panelName": "MDP",
                        "status": "success",
                        "sheetName": "MDP",
                    }
                ],
            },
        ), patch("main.webview.windows", [fake_window]), patch("main.logging.error") as log_error:
            self.api._panel_schedule_worker(job_id, payload)

        status = self.api.get_panel_schedule_background_status(job_id)
        self.assertEqual("success", status["status"])
        self.assertEqual(job_id, status["jobId"])
        self.assertEqual(1, status["panelCount"])
        self.assertEqual(1, status["completedCount"])
        self.assertEqual(1, status["successCount"])
        self.assertEqual(0, status["failureCount"])
        self.assertEqual("Added 1 panel sheet to schedule.", status["message"])
        self.assertEqual(os.path.normpath(payload["outputPath"]), status["outputPath"])
        self.assertEqual(r"C:\temp", status["outputFolder"])
        self.assertTrue(status["finishedAt"])
        self.assertFalse(self.api._panel_schedule_running)
        log_error.assert_called_once()
        self.assertIn(
            "Failed to notify Panel Schedule AI result",
            log_error.call_args[0][0],
        )

    def test_get_panel_schedule_background_status_returns_error_and_not_found_states(self):
        job_id = "job_panel_error"
        self.api._store_panel_schedule_job_record(
            self.api._build_panel_schedule_terminal_record(
                job_id,
                {
                    "status": "error",
                    "message": "Selected panel schedule was not found.",
                    "outputPath": r"C:\temp\missing.xlsx",
                    "outputFolder": r"C:\temp",
                    "successCount": 0,
                    "failureCount": 1,
                    "results": [
                        {
                            "panelId": "panel_1",
                            "panelName": "MDP",
                            "status": "error",
                            "message": "Selected panel schedule was not found.",
                        }
                    ],
                },
            )
        )

        error_status = self.api.get_panel_schedule_background_status(job_id)
        missing_status = self.api.get_panel_schedule_background_status("missing_job")

        self.assertEqual("error", error_status["status"])
        self.assertEqual(job_id, error_status["jobId"])
        self.assertEqual(1, error_status["failureCount"])
        self.assertEqual("Selected panel schedule was not found.", error_status["message"])
        self.assertEqual(
            {"status": "not_found", "jobId": "missing_job"},
            missing_status,
        )


class CanvasPanelSelectionTests(unittest.TestCase):
    """Canvas image selection -> Panel Schedule AI input mapping."""

    def setUp(self):
        self.api = Api.__new__(Api)

    def _write(self, directory, name, payload=b"binary"):
        path = Path(directory) / name
        path.write_bytes(payload)
        return str(path)

    def _stub_resolver(self, mapping):
        return lambda asset_path: mapping.get(asset_path, "")

    def _image_role(self, image_index, role, confidence="high", reason="looks like it"):
        return types.SimpleNamespace(
            image_index=image_index,
            role=role,
            confidence=confidence,
            reason=reason,
        )

    def _classification(self, panel_name, images, notes="because"):
        return types.SimpleNamespace(
            panel_name=panel_name,
            images=images,
            notes=notes,
        )

    def _entry(self, node_id, status="ok"):
        return {
            "nodeId": node_id,
            "assetPath": node_id,
            "absPath": node_id if status == "ok" else "",
            "fileName": node_id,
            "status": status,
            "message": "" if status == "ok" else "Image file not found.",
        }

    def _coerce(self, roles, extra=None):
        """Runs the coercion over one usable image per role in `roles`."""
        usable = [self._entry(f"n{i}") for i in range(len(roles))]
        resolved = usable + list(extra or [])
        parsed = self._classification(
            "MDP", [self._image_role(i, role) for i, role in enumerate(roles)]
        )
        return self.api._coerce_canvas_panel_classification(parsed, resolved, usable)

    def test_resolve_canvas_panel_selection_images_flags_missing_and_svg(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            png_path = self._write(temp_dir, "breakers.png")
            svg_path = self._write(temp_dir, "diagram.svg")
            mapping = {
                "page_assets/o/breakers.png": png_path,
                "page_assets/o/diagram.svg": svg_path,
            }
            with patch.object(
                self.api, "_resolve_page_asset_path", side_effect=self._stub_resolver(mapping)
            ):
                resolved = self.api._resolve_canvas_panel_selection_images(
                    [
                        {"nodeId": "n1", "assetPath": "page_assets/o/breakers.png"},
                        {"nodeId": "n2", "assetPath": "page_assets/o/diagram.svg"},
                        {"nodeId": "n3", "assetPath": "page_assets/o/gone.png"},
                    ]
                )

        self.assertEqual(["ok", "unsupported", "missing"], [e["status"] for e in resolved])
        self.assertEqual(png_path, resolved[0]["absPath"])
        self.assertEqual("", resolved[1]["absPath"])
        self.assertIn("SVG", resolved[1]["message"])
        self.assertEqual("Image file not found.", resolved[2]["message"])

    def test_resolve_canvas_panel_selection_images_truncates_large_selections(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            count = main.CANVAS_PANEL_CLASSIFY_MAX_IMAGES + 2
            mapping = {}
            images = []
            for index in range(count):
                asset = f"page_assets/o/img{index}.png"
                mapping[asset] = self._write(temp_dir, f"img{index}.png")
                images.append({"nodeId": f"n{index}", "assetPath": asset})
            with patch.object(
                self.api, "_resolve_page_asset_path", side_effect=self._stub_resolver(mapping)
            ):
                resolved = self.api._resolve_canvas_panel_selection_images(images)

        self.assertEqual(count, len(resolved))
        usable = [e for e in resolved if e["status"] == "ok"]
        self.assertEqual(main.CANVAS_PANEL_CLASSIFY_MAX_IMAGES, len(usable))
        self.assertIn("first", resolved[-1]["message"])

    def test_classify_canvas_panel_selection_requires_api_key(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            mapping = {"page_assets/o/a.png": self._write(temp_dir, "a.png")}
            with patch.object(
                self.api, "_resolve_page_asset_path", side_effect=self._stub_resolver(mapping)
            ), patch.object(self.api, "_resolve_panel_schedule_api_key", return_value=""):
                result = self.api.classify_canvas_panel_selection(
                    "proj_1", [{"nodeId": "n1", "assetPath": "page_assets/o/a.png"}]
                )

        self.assertEqual("error", result["status"])
        self.assertIn("GOOGLE_API_KEY", result["message"])

    def test_classify_canvas_panel_selection_returns_roles_and_abs_paths(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            breaker_path = self._write(temp_dir, "breaker.jpg")
            directory_path = self._write(temp_dir, "directory.jpg")
            mapping = {
                "page_assets/o/breaker.jpg": breaker_path,
                "page_assets/o/directory.jpg": directory_path,
            }
            parsed = self._classification(
                "  mdp  ",
                [self._image_role(0, "breaker"), self._image_role(1, "field_directory")],
            )
            with patch.object(
                self.api, "_resolve_page_asset_path", side_effect=self._stub_resolver(mapping)
            ), patch.object(
                self.api,
                "_classify_canvas_panel_images",
                side_effect=lambda usable, notes: (parsed, usable),
            ):
                result = self.api.classify_canvas_panel_selection(
                    "proj_1",
                    [
                        {"nodeId": "n1", "assetPath": "page_assets/o/breaker.jpg"},
                        {"nodeId": "n2", "assetPath": "page_assets/o/directory.jpg"},
                    ],
                    ["Panel is MDP"],
                )

        self.assertEqual("success", result["status"])
        self.assertEqual("MDP", result["panelName"])
        self.assertEqual("field_photos", result["inputMode"])
        self.assertEqual("", result["blocking"])
        self.assertEqual(
            ["breaker", "field_directory"], [i["role"] for i in result["images"]]
        )
        self.assertEqual(breaker_path, result["images"][0]["absPath"])
        self.assertEqual("high", result["images"][0]["confidence"])

    def test_a_printed_as_built_wins_over_breaker_photos(self):
        """The regression this split fixes: breakers present used to force field mode."""
        result = self._coerce(["breaker", "as_built_schedule"])

        self.assertEqual("existing_directory", result["inputMode"])
        self.assertEqual("", result["blocking"])

    def test_breakers_plus_a_field_card_read_the_breakers(self):
        result = self._coerce(["breaker", "field_directory"])

        self.assertEqual("field_photos", result["inputMode"])
        self.assertEqual("", result["blocking"])

    def test_directory_only_selections_are_transcribed(self):
        for roles in (["as_built_schedule"], ["field_directory"]):
            with self.subTest(roles=roles):
                result = self._coerce(roles)
                self.assertEqual("existing_directory", result["inputMode"])
                self.assertEqual("", result["blocking"])

    def test_classification_blocks_when_no_directory_images(self):
        result = self._coerce(["breaker"])

        self.assertEqual("field_photos", result["inputMode"])
        self.assertEqual("needs_directory", result["blocking"])
        self.assertTrue(result["blockingMessage"])

    def test_classification_blocks_when_everything_is_ignored(self):
        result = self._coerce(["ignore"])

        self.assertEqual("no_usable_images", result["blocking"])
        self.assertTrue(result["blockingMessage"])

    def test_classification_ignores_unusable_images_and_unknown_indexes(self):
        parsed = self._classification(
            "",
            [
                self._image_role(0, "breaker"),
                self._image_role(9, "as_built_schedule"),
                self._image_role(1, "not-a-role"),
            ],
        )
        usable = self._entry("n1")
        broken = self._entry("n2", status="missing")
        result = self.api._coerce_canvas_panel_classification(
            parsed, [usable, broken], [usable]
        )

        self.assertEqual("breaker", result["images"][0]["role"])
        self.assertEqual("ignore", result["images"][1]["role"])
        # The out-of-range as_built_schedule must not flip the mode.
        self.assertEqual("field_photos", result["inputMode"])
        self.assertEqual("needs_directory", result["blocking"])
        self.assertEqual("PANEL", result["panelName"])

    def test_derive_input_mode_covers_every_combination(self):
        cases = {
            (0, 0, 0): ("field_photos", "no_usable_images"),
            (1, 0, 0): ("field_photos", "needs_directory"),
            (0, 1, 0): ("existing_directory", ""),
            (0, 0, 1): ("existing_directory", ""),
            (1, 1, 0): ("existing_directory", ""),
            (1, 0, 1): ("field_photos", ""),
            (1, 1, 1): ("existing_directory", ""),
            (0, 1, 1): ("existing_directory", ""),
        }
        for (breaker, as_built, field_dir), expected in cases.items():
            with self.subTest(breaker=breaker, as_built=as_built, field_dir=field_dir):
                mode, blocking, _ = self.api._derive_canvas_panel_input_mode(
                    {
                        "breaker": breaker,
                        "as_built_schedule": as_built,
                        "field_directory": field_dir,
                    }
                )
                self.assertEqual(expected, (mode, blocking))

    def test_classify_canvas_panel_selection_errors_when_nothing_usable(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            mapping = {"page_assets/o/a.svg": self._write(temp_dir, "a.svg")}
            with patch.object(
                self.api, "_resolve_page_asset_path", side_effect=self._stub_resolver(mapping)
            ), patch.object(
                self.api,
                "_classify_canvas_panel_images",
                side_effect=AssertionError("must not call the AI"),
            ):
                result = self.api.classify_canvas_panel_selection(
                    "proj_1", [{"nodeId": "n1", "assetPath": "page_assets/o/a.svg"}]
                )

        self.assertEqual("error", result["status"])
        self.assertIn("None of the selected images could be read", result["message"])

    def test_classify_canvas_panel_selection_requires_images(self):
        result = self.api.classify_canvas_panel_selection("proj_1", [])
        self.assertEqual("error", result["status"])
        self.assertEqual("Select at least one image.", result["message"])

    def test_classification_prompt_includes_notes_and_role_definitions(self):
        prompt = self.api._build_canvas_panel_classification_prompt(3, ["Panel is MDP"])

        self.assertIn("0 through 2", prompt)
        self.assertIn("- Panel is MDP", prompt)
        self.assertIn('"breaker"', prompt)
        self.assertIn('"as_built_schedule"', prompt)
        self.assertIn('"field_directory"', prompt)

        empty = self.api._build_canvas_panel_classification_prompt(1, [])
        self.assertIn("(no notes provided)", empty)

    def test_prompt_gives_concrete_as_built_versus_field_card_discriminators(self):
        prompt = self.api._build_canvas_panel_classification_prompt(2, [])

        # Handwriting is the decisive tell, and it is stated as a hard rule.
        self.assertIn(
            'Handwriting anywhere in the circuit rows means "field_directory"', prompt
        )
        self.assertIn('NEVER an "as_built_schedule"', prompt)
        # The printed-document tells.
        self.assertIn("kVA or VA load column", prompt)
        self.assertIn("AIC", prompt)
        # The model must not be asked to pick the tool mode any more.
        self.assertNotIn("INPUT MODE", prompt)
        self.assertIn("that is worked out from the roles you assign", prompt)

    def test_normalize_canvas_panel_notes_trims_and_caps(self):
        notes = self.api._normalize_canvas_panel_notes(
            ["  spaced   out  ", "", None, "x" * 900]
        )

        self.assertEqual("spaced out", notes[0])
        self.assertEqual(main.CANVAS_PANEL_CLASSIFY_MAX_NOTE_CHARS, len(notes[1]))
        self.assertEqual(2, len(notes))

    def test_resolve_canvas_panel_schedule_output_is_stable_and_creates_nothing(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            with patch.object(
                main,
                "get_app_data_path",
                side_effect=lambda name: os.path.join(temp_dir, name),
            ):
                first = self.api.resolve_canvas_panel_schedule_output("MDP")
                second = self.api.resolve_canvas_panel_schedule_output(
                    "LP-1", first["folderToken"]
                )

            self.assertEqual("success", first["status"])
            self.assertEqual(first["outputFolder"], second["outputFolder"])
            self.assertTrue(first["outputPath"].endswith("MDP.xlsx"))
            self.assertTrue(second["outputPath"].endswith("LP-1.xlsx"))
            self.assertFalse(os.path.exists(first["outputFolder"]))

    def test_resolve_canvas_panel_schedule_output_sanitizes_the_panel_name(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            with patch.object(
                main,
                "get_app_data_path",
                side_effect=lambda name: os.path.join(temp_dir, name),
            ):
                result = self.api.resolve_canvas_panel_schedule_output('panel a/b: "main"')

        file_name = os.path.basename(result["outputPath"])
        self.assertEqual("PANEL A B MAIN.xlsx", file_name)
        self.assertNotIn("/", file_name)
        self.assertNotIn('"', file_name)

    def test_background_classification_reports_findings_through_a_job_record(self):
        finished = threading.Event()
        payload = {
            "status": "success",
            "panelName": "MDP",
            "inputMode": "field_photos",
            "blocking": "",
            "blockingMessage": "",
            "summary": "",
            "images": [],
        }
        pushed = {}

        class FakeWindow:
            def evaluate_js(self, script):
                pushed["script"] = script
                finished.set()

        with patch.object(main.webview, "windows", [FakeWindow()], create=True), patch.object(
            self.api, "classify_canvas_panel_selection", return_value=payload
        ):
            started = self.api.run_canvas_panel_classification_background(
                "proj_1", [{"nodeId": "n1", "assetPath": "page_assets/o/a.png"}]
            )
            self.assertEqual("started", started["status"])
            job_id = started["jobId"]
            self.assertTrue(finished.wait(5), "worker thread did not finish")

        record = self.api.get_canvas_panel_classification_status(job_id)
        self.assertEqual("success", record["status"])
        self.assertEqual(payload, record["result"])
        self.assertTrue(record["finishedAt"])
        self.assertIn("window.handleCanvasPanelClassificationResult(", pushed["script"])
        self.assertEqual(
            {"status": "not_found", "jobId": "missing"},
            self.api.get_canvas_panel_classification_status("missing"),
        )

    def test_background_classification_records_failures(self):
        finished = threading.Event()

        class FakeWindow:
            def evaluate_js(self, script):
                finished.set()

        with patch.object(main.webview, "windows", [FakeWindow()], create=True), patch.object(
            self.api,
            "classify_canvas_panel_selection",
            side_effect=RuntimeError("gemini exploded"),
        ):
            started = self.api.run_canvas_panel_classification_background("proj_1", [])
            job_id = started["jobId"]
            self.assertTrue(finished.wait(5), "worker thread did not finish")

        record = self.api.get_canvas_panel_classification_status(job_id)
        self.assertEqual("error", record["status"])
        self.assertEqual("gemini exploded", record["message"])
        # The guard is released so the next selection can be analyzed.
        self.assertFalse(self.api._canvas_panel_classify_job_running)

    def test_background_classification_refuses_a_second_concurrent_job(self):
        self.api._canvas_panel_classify_job_running = True
        try:
            result = self.api.run_canvas_panel_classification_background("proj_1", [])
        finally:
            self.api._canvas_panel_classify_job_running = False

        self.assertEqual("error", result["status"])
        self.assertIn("previous selection", result["message"])

    def test_classification_does_not_consume_the_shared_rate_limit(self):
        with tempfile.TemporaryDirectory(prefix="acies-canvas-panel-") as temp_dir:
            image_path = self._write(temp_dir, "a.png")
            parsed = self._classification(
                "MDP", [self._image_role(0, "as_built_schedule")]
            )
            captured = {}

            def fake_generate_content(**kwargs):
                captured["model"] = kwargs.get("model")
                return types.SimpleNamespace(parsed=parsed, text="")

            fake_client = types.SimpleNamespace(
                models=types.SimpleNamespace(generate_content=fake_generate_content)
            )

            with patch.object(main, "cb_enforce_rate_limit") as rate_limit, patch.object(
                main.genai, "Client", create=True, return_value=fake_client
            ), patch.object(
                main.types, "HttpOptions", create=True, side_effect=lambda **kw: None
            ), patch.object(
                main.types,
                "GenerateContentConfig",
                create=True,
                side_effect=lambda **kw: None,
            ), patch.object(
                main, "_open_panel_schedule_image", return_value=types.SimpleNamespace(
                    close=lambda: None
                )
            ) as open_image, patch.object(
                self.api, "_resolve_panel_schedule_api_key", return_value="key"
            ), patch.object(
                self.api, "_ensure_aiohttp", return_value=None
            ):
                usable = [
                    {
                        "nodeId": "n1",
                        "assetPath": "page_assets/o/a.png",
                        "absPath": image_path,
                        "fileName": "a.png",
                        "status": "ok",
                        "message": "",
                    }
                ]
                result, loaded = self.api._classify_canvas_panel_images(usable, [])

        rate_limit.assert_not_called()
        self.assertIs(parsed, result)
        self.assertEqual(1, len(loaded))
        self.assertEqual("gemini-3.6-flash", captured["model"])
        self.assertEqual(
            main.CANVAS_PANEL_CLASSIFY_MAX_IMAGE_EDGE,
            open_image.call_args.kwargs["max_edge"],
        )


if __name__ == "__main__":
    unittest.main()
